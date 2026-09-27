# formatter-service/tests/test_categories.py

"""
The three formatting categories, exercised through the /format API and
checked by inspecting the returned DOCX:

  - Thesis (undergraduate/masters/phd): prelim pages, roman + arabic
  - Report (report):                    no prelim pages, arabic 1..n, grammar
  - Quick Format (print_ready):         layout + arabic 1..n, no LLM at all

    python3 -m pytest tests -q
"""

import os
import sys
import tempfile
import unittest
from io import BytesIO
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

os.environ["DOCSTUDIO_PROOFREAD"] = "0"
os.environ["DOCSTUDIO_BAKE_FIELDS"] = "0"

from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app import app  # noqa: E402
from tools.inspect_docx import inspect_docx  # noqa: E402

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TYPO = "Electricty access remains low in rural areas."
FIXED = "Electricity access remains low in rural areas."


def _add_footer_field(paragraph, instr: str) -> None:
    for kind, text in (("begin", None), (None, instr), ("end", None)):
        run = paragraph.add_run()
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.text = text
        run._r.append(el)


def _report_docx() -> bytes:
    doc = Document()
    # Existing footer with text and a "Page X of Y" style field
    footer = doc.sections[0].footer
    footer.paragraphs[0].text = "University of Buea - Department of Economics"
    p = footer.add_paragraph("Page ")
    _add_footer_field(p, " PAGE ")
    p.add_run(" of ")
    _add_footer_field(p, " NUMPAGES ")

    doc.add_paragraph("A STUDY OF RURAL ELECTRIFICATION")
    doc.add_paragraph("ACKNOWLEDGEMENTS")
    doc.add_paragraph("Thanks to the Rural Electrification Agency (REA).")
    doc.add_paragraph("CHAPTER ONE: INTRODUCTION")
    doc.add_paragraph("1.1 Background of the Study")
    doc.add_paragraph(TYPO)
    doc.add_paragraph("Table 1: Households surveyed")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Households"
    doc.add_paragraph("CHAPTER TWO: LITERATURE REVIEW")
    doc.add_paragraph("2.1 Conceptual Review")
    doc.add_paragraph("Many authors discuss this topic in depth.")
    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()


class _FakeProofreader:
    """Stands in for the LLM grammar pass: fixes the one known typo."""

    calls = 0

    def correct_texts(self, texts):
        _FakeProofreader.calls += 1
        return {bid: FIXED for bid, text in texts.items() if text == TYPO}


class CategoryTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)
        _FakeProofreader.calls = 0

    def _format(self, document_type: str):
        with mock.patch.dict(os.environ, {"OPENROUTER_API_KEY": "", "DOCSTUDIO_PROOFREAD": "1"}), \
             mock.patch("formatting.formatter.Proofreader", _FakeProofreader):
            res = self.client.post(
                "/format",
                files={"file": ("in.docx", _report_docx(), DOCX_MIME)},
                data={"profileId": "ub-v1", "documentType": document_type},
            )
        self.assertEqual(res.status_code, 200, res.text)
        fd, path = tempfile.mkstemp(suffix=".docx")
        with os.fdopen(fd, "wb") as fh:
            fh.write(res.content)
        try:
            doc_text = [p.text for p in Document(path).paragraphs]
            return inspect_docx(path), doc_text
        finally:
            os.remove(path)

    def _assert_bare_page_numbers(self, report):
        for sec in report["sections"]:
            self.assertEqual(sec["footer_fields"], ["PAGE"], sec)
            self.assertEqual(sec["footer_text"], "", sec)

    def test_thesis(self):
        for level in ("undergraduate", "masters", "phd"):
            report, text = self._format(level)
            prelim, main = report["sections"]
            self.assertEqual(prelim["page_number_format"], "lowerRoman")
            self.assertEqual(main["page_number_format"], "decimal")
            self.assertEqual(main["page_number_start"], "1")
            self.assertTrue(report["has_toc_field"])
            self.assertTrue(report["has_lot_field"])
            self.assertIn("ACKNOWLEDGEMENTS", report["prelim_texts"])
            self.assertIn(FIXED, text)                       # grammar pass ran
            self.assertEqual(report["captions"][0]["text"], "Table 1.1: Households surveyed")
            self._assert_bare_page_numbers(report)

    def test_report(self):
        report, text = self._format("report")
        self.assertEqual(len(report["sections"]), 1)
        sec = report["sections"][0]
        self.assertEqual(sec["page_number_format"], "decimal")
        self.assertEqual(sec["page_number_start"], "1")
        self.assertFalse(report["has_toc_field"])
        self.assertFalse(report["has_lot_field"])
        self.assertFalse(report["has_lof_field"])
        self.assertNotIn("TABLE OF CONTENTS", text)
        self.assertNotIn("LIST OF ABBREVIATIONS", text)
        self.assertIn(FIXED, text)                           # grammar pass ran
        self.assertNotIn(TYPO, text)
        styles = {h["text"]: h["style"] for h in report["headings"]}
        self.assertEqual(styles["1.1 Background of the Study"], "Heading2")
        self.assertEqual(report["captions"][0]["text"], "Table 1.1: Households surveyed")
        self._assert_bare_page_numbers(report)

    def test_quick_format(self):
        report, text = self._format("print_ready")
        self.assertEqual(_FakeProofreader.calls, 0)          # no grammar / LLM
        self.assertIn(TYPO, text)
        self.assertEqual(len(report["sections"]), 1)
        sec = report["sections"][0]
        self.assertEqual(sec["page_number_format"], "decimal")
        self.assertEqual(sec["page_number_start"], "1")
        self.assertFalse(report["has_toc_field"])
        self.assertEqual(report["captions"], [])
        self._assert_bare_page_numbers(report)


if __name__ == "__main__":
    unittest.main()
