# formatter-service/tests/test_resilience.py

"""
Failure-path tests: LLM unavailable / erroring / returning bad JSON, and
malformed uploads. The pipeline must still produce a correctly structured
DOCX (roman prelims, arabic main) or reject the upload with a 4xx.

    python3 -m pytest tests -q
"""

import os
import sys
import tempfile
import unittest
from io import BytesIO
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

os.environ["DOCSTUDIO_PROOFREAD"] = "0"
os.environ["DOCSTUDIO_BAKE_FIELDS"] = "0"

import requests  # noqa: E402
from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app import app  # noqa: E402
from docx_parser import DOCXBlockParser  # noqa: E402
from llm.client import LLMClassifier  # noqa: E402
from llm.heuristic_classifier import classify_blocks_heuristically  # noqa: E402
from llm.schema_validator import normalize_classification, parse_llm_json  # noqa: E402
from pipeline.docx_pipeline import InvalidDocumentError, run_pipeline  # noqa: E402
from tools.inspect_docx import inspect_docx  # noqa: E402

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _report_docx() -> bytes:
    doc = Document()
    doc.add_paragraph("A STUDY OF RURAL ELECTRIFICATION")
    doc.add_paragraph("DEDICATION")
    doc.add_paragraph("To my family.")
    doc.add_paragraph("ABSTRACT")
    doc.add_paragraph("This study looks at the Rural Electrification Agency (REA).")
    doc.add_paragraph("CHAPTER ONE: INTRODUCTION")
    doc.add_paragraph("1.1 Background of the Study")
    doc.add_paragraph("Electricity access remains low in rural areas.")
    doc.add_paragraph("1.1.1 Scope")
    doc.add_paragraph("Table 1: Households surveyed")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Households"
    doc.add_paragraph("CHAPTER TWO: LITERATURE REVIEW")
    doc.add_paragraph("2.1 Conceptual Review")
    doc.add_paragraph("Many authors discuss this.")
    doc.add_paragraph("REFERENCES")
    doc.add_paragraph("Doe, J. (2020). Power for all. Energy Journal, 3(1), 1-9.")
    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()


def _write_temp(data: bytes, suffix: str = ".docx") -> str:
    fd, path = tempfile.mkstemp(suffix=suffix)
    with os.fdopen(fd, "wb") as fh:
        fh.write(data)
    return path


def _fake_llm_response(content):
    resp = mock.Mock()
    resp.raise_for_status.return_value = None
    resp.json.return_value = {"choices": [{"message": {"content": content}}]}
    return resp


class HeuristicClassifierTests(unittest.TestCase):
    def setUp(self):
        self.path = _write_temp(_report_docx())
        self.blocks = DOCXBlockParser(self.path).parse()
        self.meta = classify_blocks_heuristically(self.blocks)
        self.by_text = {
            b.get("text"): self.meta["blocks"][b["id"]]
            for b in self.blocks if b.get("type") == "paragraph"
        }

    def tearDown(self):
        os.remove(self.path)

    def test_boundary_is_first_chapter(self):
        boundary = self.meta["sections"]["prelim_ends_before_block_id"]
        text = next(b["text"] for b in self.blocks if b["id"] == boundary)
        self.assertEqual(text, "CHAPTER ONE: INTRODUCTION")

    def test_roles(self):
        self.assertEqual(self.by_text["DEDICATION"]["section"], "prelim")
        self.assertEqual(self.by_text["CHAPTER ONE: INTRODUCTION"]["role"], "chapter_heading")
        self.assertEqual(self.by_text["CHAPTER ONE: INTRODUCTION"]["chapter"], 1)
        self.assertEqual(self.by_text["CHAPTER TWO: LITERATURE REVIEW"]["chapter"], 2)
        self.assertEqual(self.by_text["1.1 Background of the Study"]["headingLevel"], 2)
        self.assertEqual(self.by_text["1.1.1 Scope"]["headingLevel"], 3)
        self.assertEqual(self.by_text["REFERENCES"]["role"], "references_heading")
        self.assertEqual(self.by_text["Table 1: Households surveyed"]["role"], "caption")

    def test_table_linked_to_caption(self):
        table_meta = self.meta["blocks"]["T1"]
        caption_id = table_meta["caption_block_id"]
        caption = next(b["text"] for b in self.blocks if b["id"] == caption_id)
        self.assertEqual(caption, "Table 1: Households surveyed")
        self.assertEqual(table_meta["chapter"], 1)


class LLMJsonTests(unittest.TestCase):
    blocks = [
        {"id": "P1", "type": "paragraph", "text": "Front"},
        {"id": "P2", "type": "paragraph", "text": "CHAPTER ONE"},
    ]

    def test_parse_tolerates_fences_and_prose(self):
        raw = 'Sure! Here it is:\n```json\n{"blocks": {"P1": {"role": "body_paragraph"}}}\n```\nDone.'
        self.assertIn("blocks", parse_llm_json(raw))

    def test_parse_rejects_garbage(self):
        for raw in ("", "no json here", "[1, 2, 3]", "{broken json", None):
            with self.assertRaises(ValueError):
                parse_llm_json(raw)

    def test_normalize_rejects_wrong_shape(self):
        for data in ({}, {"blocks": []}, {"blocks": {"X9": {"role": "body_paragraph"}}},
                     {"blocks": {"P1": "chapter_heading"}}):
            with self.assertRaises(ValueError):
                normalize_classification(data, self.blocks)

    def test_normalize_sets_media_chapter_from_position(self):
        blocks = [
            {"id": "P1", "type": "paragraph", "text": "CHAPTER ONE"},
            {"id": "T1", "type": "table"},
            {"id": "P2", "type": "paragraph", "text": "CHAPTER TWO"},
            {"id": "P3", "type": "paragraph", "text": ""},
            {"id": "F1", "type": "image", "parent": "P3"},
        ]
        data = {"blocks": {
            "P1": {"role": "chapter_heading", "chapter": 1},
            "T1": {"role": "table", "chapter": 2},        # wrong
            "P2": {"role": "chapter_heading"},            # chapter omitted
            "F1": {"role": "figure"},                     # chapter omitted
        }}
        out = normalize_classification(data, blocks)
        self.assertEqual(out["blocks"]["T1"]["chapter"], 1)
        self.assertEqual(out["blocks"]["P2"]["chapter"], 2)
        self.assertEqual(out["blocks"]["F1"]["chapter"], 2)

    def test_normalize_repairs_bad_boundary(self):
        fallback = classify_blocks_heuristically(self.blocks)
        data = {"blocks": {"P1": {"role": "body_paragraph", "section": "bogus"}},
                "sections": {"prelim_ends_before_block_id": "P999"}}
        out = normalize_classification(data, self.blocks, fallback=fallback)
        self.assertEqual(out["sections"]["prelim_ends_before_block_id"], "P2")
        self.assertEqual(out["blocks"]["P1"]["section"], "prelim")
        self.assertEqual(out["blocks"]["P2"]["role"], "chapter_heading")


class LLMFailureTests(unittest.TestCase):
    def setUp(self):
        self.path = _write_temp(_report_docx())
        self.blocks = DOCXBlockParser(self.path).parse()

    def tearDown(self):
        os.remove(self.path)

    def _classify(self, **post_kwargs):
        with mock.patch("llm.client.requests.post", **post_kwargs) as post:
            meta = LLMClassifier(api_key="test-key").classify_document_blocks(self.blocks)
        return meta, post

    def test_no_api_key_uses_heuristic(self):
        with mock.patch.dict(os.environ, {"OPENROUTER_API_KEY": ""}):
            meta = LLMClassifier().classify_document_blocks(self.blocks)
        self.assertEqual(meta["classifier"], "heuristic")

    def test_network_error_falls_back(self):
        meta, post = self._classify(side_effect=requests.ConnectionError("down"))
        self.assertEqual(meta["classifier"], "heuristic")
        self.assertGreaterEqual(post.call_count, 1)

    def test_http_error_falls_back(self):
        resp = mock.Mock()
        resp.raise_for_status.side_effect = requests.HTTPError("429 Too Many Requests")
        meta, _ = self._classify(return_value=resp)
        self.assertEqual(meta["classifier"], "heuristic")

    def test_invalid_json_falls_back(self):
        meta, post = self._classify(return_value=_fake_llm_response("I cannot help with that."))
        self.assertEqual(meta["classifier"], "heuristic")
        self.assertEqual(post.call_count, 3)  # tried every model

    def test_unexpected_response_shape_falls_back(self):
        resp = mock.Mock()
        resp.raise_for_status.return_value = None
        resp.json.return_value = {"error": "oops"}
        meta, _ = self._classify(return_value=resp)
        self.assertEqual(meta["classifier"], "heuristic")

    def test_valid_llm_json_is_used(self):
        content = '{"blocks": {"P1": {"role": "title_page", "section": "prelim"}}, "sections": {}}'
        meta, _ = self._classify(return_value=_fake_llm_response(content))
        self.assertTrue(meta["classifier"].startswith("llm:"))
        self.assertEqual(meta["blocks"]["P1"]["role"], "title_page")
        # gaps filled from the heuristic, including the boundary
        boundary = meta["sections"]["prelim_ends_before_block_id"]
        self.assertEqual(meta["blocks"][boundary]["role"], "chapter_heading")

    def test_pipeline_output_with_llm_failure(self):
        with mock.patch("llm.client.requests.post", side_effect=requests.Timeout("slow")), \
             mock.patch.dict(os.environ, {"OPENROUTER_API_KEY": "test-key"}):
            out = run_pipeline(self.path, "default")
        out_path = _write_temp(out)
        try:
            report = inspect_docx(out_path)
        finally:
            os.remove(out_path)

        self.assertEqual(len(report["sections"]), 2)
        prelim, main = report["sections"]
        self.assertEqual(prelim["page_number_format"], "lowerRoman")
        self.assertEqual(prelim["page_number_start"], "1")
        self.assertIn("PAGE", prelim["footer_fields"])
        self.assertEqual(main["page_number_format"], "decimal")
        self.assertEqual(main["page_number_start"], "1")
        self.assertIn("PAGE", main["footer_fields"])
        self.assertTrue(main["first_text"].upper().startswith("CHAPTER 1"))
        self.assertIn("DEDICATION", report["prelim_texts"])
        self.assertTrue(report["has_toc_field"])
        self.assertTrue(report["update_fields_on_open"])
        styles = {h["text"]: h["style"] for h in report["headings"]}
        self.assertEqual(styles["1.1 Background of the Study"], "Heading2")
        self.assertEqual(styles["1.1.1 Scope"], "Heading3")
        self.assertEqual(report["tables"], 1)


class MalformedDocxTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    def _post(self, data: bytes, name="upload.docx", document_type="report"):
        return self.client.post(
            "/format",
            files={"file": (name, data, DOCX_MIME)},
            data={"profileId": "default", "documentType": document_type},
        )

    def test_pipeline_rejects_non_zip(self):
        path = _write_temp(b"%PDF-1.4 not a docx")
        try:
            with self.assertRaises(InvalidDocumentError):
                run_pipeline(path, "default")
        finally:
            os.remove(path)

    def test_api_rejects_garbage(self):
        for mode in ("report", "print_ready"):
            res = self._post(b"this is plainly not a word document", document_type=mode)
            self.assertEqual(res.status_code, 422, res.text)
            self.assertIn("not a valid .docx", res.json()["detail"])

    def test_api_rejects_empty(self):
        res = self._post(b"")
        self.assertEqual(res.status_code, 422)

    def test_api_rejects_zip_without_document(self):
        import zipfile
        bio = BytesIO()
        with zipfile.ZipFile(bio, "w") as zf:
            zf.writestr("hello.txt", "hi")
        res = self._post(bio.getvalue())
        self.assertEqual(res.status_code, 422)
        self.assertIn("not a Word document", res.json()["detail"])

    def test_api_rejects_truncated_docx(self):
        res = self._post(_report_docx()[:2000])
        self.assertEqual(res.status_code, 422, res.text)

    def test_api_rejects_document_without_text(self):
        bio = BytesIO()
        Document().save(bio)
        res = self._post(bio.getvalue())
        self.assertEqual(res.status_code, 422)

    def test_api_formats_valid_docx(self):
        with mock.patch.dict(os.environ, {"OPENROUTER_API_KEY": ""}):
            res = self._post(_report_docx())
        self.assertEqual(res.status_code, 200, res.text)
        self.assertEqual(res.content[:2], b"PK")


if __name__ == "__main__":
    unittest.main()
